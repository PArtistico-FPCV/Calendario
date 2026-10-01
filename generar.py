# Lee calendario-fpcv-eventos.xlsx y genera _site/index.html a partir de plantilla.html
import json, os, re, sys
from datetime import date, datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

XLSX = "calendario-fpcv-eventos.xlsx"
CAT = {
    "competición fpcv":"comp","competicion fpcv":"comp",
    "campeonato de españa":"esp","campeonato de espana":"esp",
    "competición internacional":"int","competicion internacional":"int",
    "trofeo club":"trofeo","trofeo de club":"trofeo",
    "curso":"curso",
    "reunión":"reun","reunion":"reun",
    "prueba de acceso":"acceso",
}
HOR = {"mañana":"M","manana":"M","tarde":"T","todo el día":"","todo el dia":""}
DEFAULT_N = {"comp":1,"esp":1,"int":1,"trofeo":1,"acceso":1,"curso":0,"reun":0}  # si Huecos está vacío

def s(v): return str(v).strip() if v is not None else ""
def as_date(v):
    if isinstance(v,datetime): return v.date()
    return v if isinstance(v,date) else None

wb = load_workbook(XLSX, data_only=True)
ws = wb["Eventos"]
ev, errors = [], []
for i,row in enumerate(ws.iter_rows(min_row=2, max_col=9, values_only=True), start=2):
    a,b,t,p,c,n,hv,hs,hd = (list(row)+[None]*9)[:9]
    if all(x in (None,"") for x in (a,b,t,p,c,n)): continue
    a, b = as_date(a), as_date(b)
    if not a: errors.append(f"Fila {i}: falta la fecha de Inicio o no es una fecha válida"); continue
    if b is None: b = a
    if b < a: errors.append(f"Fila {i}: la fecha de Fin es anterior al Inicio"); continue
    if not s(t): errors.append(f"Fila {i}: falta el nombre del Evento"); continue
    k = CAT.get(s(c).lower())
    if not k: errors.append(f"Fila {i}: Categoría no reconocida («{s(c)}»). Elige una de la lista"); continue
    if s(n) == "": huecos = DEFAULT_N[k]
    else:
        try: huecos = int(float(s(n).replace(",",".")))
        except ValueError: huecos = -1
        if huecos not in (0,1,2): errors.append(f"Fila {i}: Huecos tiene que ser 0, 1 o 2 (pone «{s(n)}»)"); continue
    days = [(a.toordinal()+x)%7 for x in range((b-a).days+1)]   # 5=vie, 6=sáb, 0=dom
    if not any(x in (5,6,0) for x in days):
        print(f"AVISO fila {i}: «{s(t)}» no cae en viernes, sábado o domingo y no se mostrará")
    h = {}
    for dow,val in ((5,hv),(6,hs),(0,hd)):
        code = HOR.get(s(val).lower())
        if s(val) and code is None: errors.append(f"Fila {i}: horario no válido («{s(val)}»). Usa Mañana, Tarde o Todo el día")
        elif code: h[str(dow)] = code
    item = {"s":a.isoformat(),"e":b.isoformat(),"t":s(t),"p":s(p),"c":k,"n":huecos}
    if h: item["h"] = h
    ev.append(item)

if errors:
    print("\nNo se ha publicado porque hay errores en el Excel:\n")
    for e in errors: print(" -", e)
    sys.exit(1)
if not ev:
    print("El Excel no tiene ningún evento."); sys.exit(1)

ev.sort(key=lambda e:(e["s"],e["e"],e["t"]))
html = open("plantilla.html",encoding="utf-8").read()
hoy = datetime.now(ZoneInfo("Europe/Madrid")).strftime("%d/%m/%Y")
datos = json.dumps(ev, ensure_ascii=False, indent=2)
html = re.sub(r"/\*DATA_START\*/.*?/\*DATA_END\*/", lambda m: "/*DATA_START*/"+datos+"/*DATA_END*/", html, flags=re.S)
html = html.replace("__UPDATED__", hoy)
os.makedirs("_site", exist_ok=True)
open("_site/index.html","w",encoding="utf-8").write(html)
print(f"OK: {len(ev)} eventos -> _site/index.html")
